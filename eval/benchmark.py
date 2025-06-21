import sys
sys.path.append(".")

from tqdm import tqdm
import argparse
import pandas as pd
import ast

import torch

from utils import eval_backup, utils
from utils.eval_backup import _segment_net_transition
from utils.dataset import MotionDataset
from model.segment_net import SegmentNet
from model.twostage import ContextTransformer

def arg_as_list(s):
    try:
        v = ast.literal_eval(s)
        if type(v) is not list:
            raise argparse.ArgumentTypeError("Argument must be a list, e.g., '[90]' or '[15, 60]'")
        return v
    except (ValueError, SyntaxError):
        raise argparse.ArgumentTypeError("Argument must be a Python-style list, e.g., '[90]' or '[15, 60]'")

if __name__ =="__main__":
    # arguments
    parser = argparse.ArgumentParser()
    parser.add_argument("--dataset", type=str, required=True)
    parser.add_argument("--interp", type=lambda s: s.lower() in ['true', '1'])
    parser.add_argument("--config", type=str, default="default.yaml")
    parser.add_argument("--ts_configs", type=str, nargs="+", default=[])
    parser.add_argument("--rmi_configs", type=str, nargs="+", default=[])
    parser.add_argument("--ours_configs", type=str, nargs="+", default=[])
    parser.add_argument("--traj_edit", type=str, default=None)
    parser.add_argument("--intra_batch_replace", action="store_true")
    parser.add_argument("--replace_target", action="store_true")
    parser.add_argument("--multi_ib", action="store_true")
    parser.add_argument("--use_segment", action="store_true")
    parser.add_argument("--only_segment", action="store_true")
    parser.add_argument("--kf_config", type=str, default="keyframe.yaml")
    parser.add_argument("--seg_config", type=str, default="segment.yaml")
    parser.add_argument("--kf_sampling", type=str, default="score")
    parser.add_argument("--kf_param", type=float, default=None)
    parser.add_argument("--gpu", type=int, default=0, help="GPU device ID")
    parser.add_argument("--shuffle", action="store_true", help="Enable data shuffling. By default, no shuffle is used.")

    # Arguments for different evaluation modes
    parser.add_argument("--eval_specific_seqs", action="store_true", help="Evaluate a specific range of sequences.")
    parser.add_argument("--batch_idx", type=int, default=0, help="The index of the batch to evaluate.")
    parser.add_argument("--num_seqs", type=int, default=1, help="Number of sequences to evaluate within the batch. Use -1 for all.")
    parser.add_argument("--eval_all_seqs_to_excel", action="store_true", help="Evaluate all sequences in all batches and save to a multi-sheet Excel file.")
    parser.add_argument("--output_file", type=str, default="evaluation_results.xlsx", help="Output file for the evaluation results (.csv or .xlsx).")
    parser.add_argument("--transition_lengths", type=arg_as_list, default=None, help="A list of specific transition lengths to evaluate, e.g., '[90]' or '[15, 60]'. Defaults to all.")


    args = parser.parse_args()

    # Set device
    device = torch.device(f"cuda:{args.gpu}" if torch.cuda.is_available() else "cpu")

    # Define the transition lengths to be evaluated
    if args.transition_lengths:
        transitions = args.transition_lengths
        print(f"Using specific transition lengths: {transitions}")
    else:
        transitions = [15, 30, 60, 90]
        print(f"Using default transition lengths: {transitions}")

    # Logic for evaluating a specific range of sequences and saving to Excel
    if args.eval_specific_seqs or args.eval_all_seqs_to_excel:
        try:
            import openpyxl
        except ImportError:
            print("Error: 'openpyxl' is required to write Excel files. Please install it using 'pip install openpyxl'")
            sys.exit(1)
        
        if args.eval_all_seqs_to_excel:
            print("--- Full Dataset Evaluation to Excel Mode ---")
        else:
            print("--- Specific Sequence Evaluation Mode ---")

        evaluator = eval_backup.Evaluator(args)
        skeleton = evaluator.skeleton
        results_table = []

        # The outer loop must be the transition lengths because the evaluator's generator is consumed on each full run.
        for trans in tqdm(transitions, desc="Processing transition lengths"):
            num_frames = evaluator.config.context_frames + trans + 1
            
            # The evaluator creates a new dataloader each time eval() is called.
            # It will iterate through all batches for the given num_frames.
            eval_generator = evaluator.eval(num_frames, traj_option=args.traj_edit)
            
            for batch_idx, res in enumerate(tqdm(eval_generator, desc=f"Batch (trans={trans})", leave=False)):
                # If in specific mode, skip until the target batch is found
                if args.eval_specific_seqs and not args.eval_all_seqs_to_excel and batch_idx != args.batch_idx:
                    continue

                batch_size = res["motions"][0].shape[0]
                
                # Determine number of sequences to process in this batch
                if args.eval_all_seqs_to_excel:
                    num_seqs_to_eval = batch_size
                else: # eval_specific_seqs mode
                    num_seqs_to_eval = args.num_seqs
                    if num_seqs_to_eval == -1 or num_seqs_to_eval > batch_size:
                        num_seqs_to_eval = batch_size

                # Loop through each sequence in the current batch
                for seq_idx in range(num_seqs_to_eval):
                    gt_motion_single = res["motions"][0][seq_idx:seq_idx+1]
                    gt_contact_single = res["contacts"][0][seq_idx:seq_idx+1]
                    
                    # Loop through each method for the sequence
                    num_methods = len(res["tags"])
                    for method_idx in range(1, num_methods): # Start from 1 to skip GT
                        method_tag = res["tags"][method_idx]
                        pred_motion_single = res["motions"][method_idx][seq_idx:seq_idx+1]
                        pred_contact_single = res["contacts"][method_idx][seq_idx:seq_idx+1]

                        # Calculate metrics
                        l2p_val = eval_backup.l2p(gt_motion_single, pred_motion_single, skeleton, evaluator.l2p_mean, evaluator.l2p_std, evaluator.config.context_frames)
                        l2q_val = eval_backup.l2q(gt_motion_single, pred_motion_single, evaluator.config.context_frames)
                        npss_val = eval_backup.npss(gt_motion_single, pred_motion_single, evaluator.config.context_frames)
                        fs_val = eval_backup.foot_skate(pred_motion_single, pred_contact_single, skeleton, evaluator.contact_idx, ctx_frames=evaluator.config.context_frames)

                        row = {
                            "Method": method_tag,
                            "Batch": batch_idx,
                            "Sequence": seq_idx,
                            "Transition Length": trans,
                            "L2P": l2p_val,
                            "L2Q": l2q_val,
                            "NPSS": npss_val * 10,
                            "Foot Skate": fs_val * 10,
                        }
                        results_table.append(row)
                
                # If only evaluating one batch, break after processing it
                if args.eval_specific_seqs and not args.eval_all_seqs_to_excel:
                    break

        # After all loops, create the Excel file with multiple sheets
        print("All evaluations complete. Saving results to Excel file...")
        df = pd.DataFrame(results_table)

        with pd.ExcelWriter(args.output_file, engine='openpyxl') as writer:
            unique_transitions = sorted(df['Transition Length'].unique())
            for trans_len in unique_transitions:
                df_sheet = df[df['Transition Length'] == trans_len].copy()
                
                # Drop the now-redundant 'Transition Length' column
                df_sheet.drop('Transition Length', axis=1, inplace=True)
                
                sheet_name = f'Transition_{trans_len}'
                df_sheet.to_excel(writer, sheet_name=sheet_name, index=False, float_format="%.4f")
                print(f"  - Sheet '{sheet_name}' created.")

        print(f"\nResults saved to {args.output_file}")

    else:
        # Original logic for evaluating the entire dataset and printing to a .txt file
        if args.kf_param is not None: kf_sampling = [args.kf_sampling, args.kf_param]
        else: kf_sampling = [args.kf_sampling]
        
        evaluator = eval_backup.Evaluator(args)
        results = { "tags": None, "l2p": [], "l2q": [], "npss": [], "foot skate": [] }
        
        segment_model, keyframe_model, dataset = None, None, None
        if results["tags"] is None: results["tags"] = []
        
        if args.use_segment:
            print("Loading SegmentNet model...")
            dataset = MotionDataset(train=False, config=evaluator.config)
            mean, std = dataset.motion_statistics(device)
            traj_mean, traj_std = dataset.traj_statistics(device)
            kf_config = utils.load_config(f"config/{args.dataset}/{args.kf_config}")
            keyframe_model = ContextTransformer(kf_config, dataset).to(device)
            utils.load_model(keyframe_model, kf_config)
            keyframe_model.eval()
            seg_config = utils.load_config(f"config/{args.dataset}/{args.seg_config}")
            segment_model = SegmentNet(seg_config, dataset).to(device)
            utils.load_model(segment_model, seg_config)
            segment_model.eval()
            if "SegmentNet" not in results["tags"]: results["tags"].append("SegmentNet")

        for trans in tqdm(transitions):
            num_frames = evaluator.config.context_frames + trans + 1
            motion_list, contact_list, traj_list = [], [], []
            skeleton = None
            
            baseline_results = []
            for idx, res in enumerate(evaluator.eval(num_frames, traj_option=args.traj_edit)):
                baseline_results.append(res)
                if idx == 0:
                    skeleton = res["skeleton"]
                    for tag in res["tags"][1:]:
                        if tag not in results["tags"]: results["tags"].append(tag)
            
            for res in baseline_results:
                motion_list.append(res["motions"])
                contact_list.append(res["contacts"])
                traj_list.append(res["trajs"])
            
            if args.use_segment and segment_model:
                GT_motion, GT_contact, GT_traj = baseline_results[0]["motions"][0], baseline_results[0]["contacts"][0], baseline_results[0]["trajs"][0]
                GT_phase, GT_score = None, None
                if seg_config.use_phase:
                    GT_phase = next(iter(torch.utils.data.DataLoader(dataset, batch_size=GT_motion.shape[0])))["phase"].to(device)[:, :num_frames]
                if seg_config.use_score:
                    GT_score = next(iter(torch.utils.data.DataLoader(dataset, batch_size=GT_motion.shape[0])))["score"].to(device)[:, :num_frames]
                
                with torch.no_grad():
                    segment_result = _segment_net_transition(seg_config, keyframe_model, segment_model, GT_motion, mean, std, GT_contact, GT_phase, GT_traj, GT_score, traj_mean, traj_std, kf_sampling)
                
                seg_motion, seg_contact = segment_result["motion"], segment_result.get("contact", GT_contact)
                # Shape matching logic
                if seg_motion.shape[1] != GT_motion.shape[1]:
                    padding = torch.zeros((seg_motion.shape[0], GT_motion.shape[1] - seg_motion.shape[1], seg_motion.shape[2]), device=device)
                    seg_motion = torch.cat([seg_motion, padding], dim=1) if seg_motion.shape[1] < GT_motion.shape[1] else seg_motion[:, :GT_motion.shape[1]]
                if seg_contact.shape[1] != GT_contact.shape[1]:
                    padding = torch.zeros((seg_contact.shape[0], GT_contact.shape[1] - seg_contact.shape[1], seg_contact.shape[2]), device=device)
                    seg_contact = torch.cat([seg_contact, padding], dim=1) if seg_contact.shape[1] < GT_contact.shape[1] else seg_contact[:, :GT_contact.shape[1]]

                motion_list.append([GT_motion, seg_motion])
                contact_list.append([GT_contact, seg_contact])
                traj_list.append([GT_traj, GT_traj])
            
            final_motion_list = [torch.cat([b[i] for b in motion_list], dim=0) for i in range(len(motion_list[0]))]
            final_contact_list = [torch.cat([b[i] for b in contact_list], dim=0) for i in range(len(contact_list[0]))]

            results["l2p"].append([eval_backup.l2p(final_motion_list[0], m, skeleton, evaluator.l2p_mean, evaluator.l2p_std) for m in final_motion_list[1:]])
            results["l2q"].append([eval_backup.l2q(final_motion_list[0], m) for m in final_motion_list[1:]])
            results["npss"].append([eval_backup.npss(final_motion_list[0], m) for m in final_motion_list[1:]])
            results["foot skate"].append([eval_backup.foot_skate(m, c, skeleton, evaluator.contact_idx) for m, c in zip(final_motion_list[1:], final_contact_list[1:])])
        
        # Original text file saving logic
        output_filename = f"eval/benchmark-{args.dataset}{'-segment' if args.use_segment else ''}.txt"
        with open(output_filename, "w") as f:
            metric_list = ["l2p", "l2q", "npss", "foot skate"]
            if args.traj_edit is not None: metric_list.append("traj")
            for metric in metric_list:
                f.write(f"{metric.upper() if metric != 'foot skate' else 'Foot skate'}\n")
                row = f"$t_\\mathrm{{trans}}$"
                for t in transitions: row += f" & {t}"
                f.write(f"{row} \\\\ \\midrule\n")
                for i in range(len(results["tags"])): 
                    method = results["tags"][i]
                    row_str = str(method)
                    for j in range(len(transitions)):
                        row_str += f" & {results[metric][j][i]:.2f}" if metric in ["l2p", "l2q"] else f" & {results[metric][j][i] * 10:.2f}"
                    f.write(row_str + " \\\\\n")
                f.write("\n")
        
        print("Benchmark evaluation finished.")